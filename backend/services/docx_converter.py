import io
import logging
from typing import Optional
import tempfile
import os

logger = logging.getLogger(__name__)

class DocxConverter:
    """DOCX dosyalarını işlemek için yardımcı sınıf"""
    
    @staticmethod
    def validate_docx(document_bytes: bytes) -> bool:
        """
        DOCX dosyasının geçerli olup olmadığını kontrol eder
        
        Args:
            document_bytes: DOCX dosyası bytes formatında
            
        Returns:
            bool: Dosya geçerliyse True
        """
        try:
            # DOCX dosyalarının ZIP yapısını kontrol et
            import zipfile
            
            with io.BytesIO(document_bytes) as docx_stream:
                with zipfile.ZipFile(docx_stream, 'r') as zip_file:
                    # DOCX dosyalarında olması gereken temel dosyalar
                    required_files = ['[Content_Types].xml', 'word/document.xml']
                    zip_contents = zip_file.namelist()
                    
                    for required_file in required_files:
                        if required_file not in zip_contents:
                            logger.error(f"DOCX dosyasında eksik dosya: {required_file}")
                            return False
                    
                    logger.info("DOCX dosyası geçerli ZIP yapısına sahip")
                    return True
                    
        except Exception as e:
            logger.error(f"DOCX validasyonu sırasında hata: {str(e)}")
            return False
    
    @staticmethod
    def extract_text_from_docx(document_bytes: bytes) -> Optional[str]:
        """
        DOCX dosyasından metin çıkarır (fallback seçeneği olarak)
        
        Args:
            document_bytes: DOCX dosyası bytes formatında
            
        Returns:
            str: Çıkarılan metin veya None
        """
        try:
            from docx import Document
            
            with io.BytesIO(document_bytes) as docx_stream:
                doc = Document(docx_stream)
                
                # Tüm paragrafları al
                text_content = []
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        text_content.append(paragraph.text.strip())
                
                # Tabloları da al
                for table in doc.tables:
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            if cell.text.strip():
                                row_text.append(cell.text.strip())
                        if row_text:
                            text_content.append(" | ".join(row_text))
                
                full_text = "\n".join(text_content)
                logger.info(f"DOCX'den {len(full_text)} karakter metin çıkarıldı")
                return full_text
                
        except ImportError:
            logger.warning("python-docx kütüphanesi bulunamadı")
            return None
        except Exception as e:
            logger.error(f"DOCX metin çıkarma sırasında hata: {str(e)}")
            return None
    
    @staticmethod
    def create_temp_file(document_bytes: bytes, suffix: str = ".docx") -> Optional[str]:
        """
        Geçici dosya oluşturur
        
        Args:
            document_bytes: Dosya içeriği
            suffix: Dosya uzantısı
            
        Returns:
            str: Geçici dosya yolu veya None
        """
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
                temp_file.write(document_bytes)
                temp_path = temp_file.name
                
            logger.info(f"Geçici dosya oluşturuldu: {temp_path}")
            return temp_path
            
        except Exception as e:
            logger.error(f"Geçici dosya oluşturma sırasında hata: {str(e)}")
            return None
    
    @staticmethod
    def cleanup_temp_file(file_path: str):
        """Geçici dosyayı siler"""
        try:
            if os.path.exists(file_path):
                os.unlink(file_path)
                logger.info(f"Geçici dosya silindi: {file_path}")
        except Exception as e:
            logger.error(f"Geçici dosya silme sırasında hata: {str(e)}")
